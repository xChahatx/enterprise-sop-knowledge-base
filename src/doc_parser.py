"""
Document Parser Module
Extracts text from PDF files and prepares metadata
for versioning, access control and semantic search.
"""

import fitz
import pdfplumber
from docx import Document
import hashlib
from pathlib import Path
from typing import Tuple
import logging

logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)


class DocumentParser:

    def __init__(self):
        self.supported_formats = [".pdf",
                                  ".docx"
        ]

    # ------------------------------------------------------------------
    # TABLE EXTRACTION
    # ------------------------------------------------------------------

    def extract_tables(self, pdf_path: str) -> str:

        table_text = []

        try:

            with pdfplumber.open(pdf_path) as pdf:

                for page_no, page in enumerate(pdf.pages):

                    tables = page.extract_tables()

                    if not tables:
                        continue

                    table_text.append(
                        f"\n\n===== TABLES FROM PAGE {page_no + 1} =====\n"
                    )

                    for table_number, table in enumerate(tables, start=1):

                        if not table or len(table) < 2:
                            continue

                        table_text.append(
                            f"Table {table_number}"
                        )

                        headers = []

                        first_row = table[0]

                        if first_row:

                            headers = [
                                str(cell).strip() if cell else f"Column {i+1}"
                                for i, cell in enumerate(first_row)
                            ]

                        for row_number, row in enumerate(table[1:], start=1):

                            if not row:
                                continue

                            row_lines = []

                            for col_index, cell in enumerate(row):

                                header = (
                                    headers[col_index]
                                    if col_index < len(headers)
                                    else f"Column {col_index+1}"
                                )

                                value = ""

                                if cell is not None:
                                    value = str(cell).replace("\n", " ").strip()

                                if value:
                                    row_lines.append(
                                        f"{header}: {value}"
                                    )

                            if row_lines:

                                table_text.append(
                                    f"Row {row_number}"
                                )

                                table_text.extend(row_lines)

                                table_text.append("")

        except Exception as e:
            logger.warning(f"Table extraction failed: {e}")

        return "\n".join(table_text)
    def parse_docx(self, file_path: str) -> Tuple[str, dict]:
        file_path = Path(file_path)
        file_hash = hashlib.md5(
            file_path.read_bytes()
        ).hexdigest()
        document = Document(file_path)
        
        content = []
        for para in document.paragraphs:
            text = para.text.strip()
            if not text:
                continue
            style = para.style.name.lower()
            if "heading" in style:
                content.append(f"# {text}")
            elif "bullet" in style:
                content.append(f"• {text}")
            elif "number" in style or "list" in style:
                content.append(text)
            else:
                content.append(text)
        for table_index, table in enumerate(document.tables, start=1):
            content.append("")
            content.append(f"===== TABLE {table_index} =====")
            if not table.rows:
                continue
            
            headers = []
            for cell in table.rows[0].cells:
                headers.append(cell.text.strip())
            for row_number, row in enumerate(table.rows[1:], start=1):
                content.append(f"Row {row_number}")
                for i, cell in enumerate(row.cells):
                    header = headers[i] if i < len(headers) else f"Column {i+1}"
                    value = cell.text.strip()
                    if value:
                        content.append(
                            f"{header}: {value}"
                        )
                content.append("")


        extracted_text = "\n".join(content)
        metadata = {
            "filename": file_path.name,
            "file_hash": file_hash,
            "file_extension": ".docx",
            "file_size_kb": round(
                file_path.stat().st_size / 1024,
                2
            ),
            "num_pages": 1,
            "total_characters": len(extracted_text),
            "document_type": "SOP",
            "language": "English",
            "parser": "python-docx",
            "contains_tables": len(document.tables) > 0,
            "page_map": [
                {
                    "page": 1,
                    "text": extracted_text
                }
            ]
        }
        logger.info(
            f"Successfully parsed DOCX: {file_path.name}"
        )
        logger.info(
            f"Extracted {metadata['total_characters']} characters"
        )

        return extracted_text, metadata
    # ------------------------------------------------------------------
    # PDF PARSER
    # ------------------------------------------------------------------

    def parse_pdf(self, file_path: str) -> Tuple[str, dict]:

        try:

            file_path = Path(file_path)

            if not file_path.exists():
                raise FileNotFoundError(
                    f"File not found: {file_path}"
                )

            file_hash = hashlib.md5(
                file_path.read_bytes()
            ).hexdigest()

            pdf_doc = fitz.open(file_path)

            full_text = []
            page_map = []

            for page_num in range(len(pdf_doc)):

                page = pdf_doc[page_num]

                text = page.get_text()

                if text.strip():

                    full_text.append(text)

                    page_map.append(
                        {
                            "page": page_num + 1,
                            "text": text
                        }
                    )

            pdf_doc.close()

            extracted_text = "\n".join(full_text)

            table_text = self.extract_tables(
                str(file_path)
            )

            if table_text.strip():

                extracted_text += "\n\n"

                extracted_text += table_text

            metadata = {

                "filename": file_path.name,
                "file_hash": file_hash,
                "file_extension": file_path.suffix.lower(),
                "file_size_kb": round(
                    file_path.stat().st_size / 1024,
                    2
                ),

                "num_pages": len(page_map),
                "total_characters": len(extracted_text),
                "document_type": "SOP",
                "language": "English",
                "parser": "PyMuPDF + pdfplumber",
                "contains_tables": bool(table_text.strip()),
                "page_map": page_map
            }

            logger.info(
                f"Successfully parsed PDF: {file_path.name}"
            )

            logger.info(
                f"Extracted {metadata['num_pages']} pages | "
                f"{metadata['total_characters']} characters"
            )

            return extracted_text, metadata

        except Exception as e:

            logger.error(f"Error parsing PDF: {e}")

            raise


    # ------------------------------------------------------------------
    # PARSE SINGLE DOCUMENT
    # ------------------------------------------------------------------

    def parse(self, file_path: str) -> Tuple[str, dict]:

        file_path = Path(file_path)

        if not file_path.exists():
            raise FileNotFoundError(
                f"File not found: {file_path}"
            )

        suffix = file_path.suffix.lower()

        if suffix not in self.supported_formats:
            raise ValueError(
                f"Unsupported format: {suffix}"
            )

        if suffix == ".pdf":
            return self.parse_pdf(str(file_path))
        elif suffix == ".docx":
            return self.parse_docx(str(file_path))

    # ------------------------------------------------------------------
    # PARSE MULTIPLE DOCUMENTS
    # ------------------------------------------------------------------

    def parse_multiple(self, folder_path: str):

        folder = Path(folder_path)

        if not folder.exists():
            raise FileNotFoundError(folder_path)

        results = []

        files = []

        for ext in self.supported_formats:
            files.extend(folder.glob(f"*{ext}"))

        logger.info(
            f"Found {len(files)} documents"
        )

        for file in files:

            try:

                text, metadata = self.parse(
                    str(file)
                )

                results.append(
                    (
                        text,
                        metadata
                    )
                )

            except Exception as e:

                logger.warning(
                    f"Skipped {file.name}: {e}"
                )

        return results